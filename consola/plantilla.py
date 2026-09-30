"""La plantilla del inventario: el .xlsx que se llena y se sube en "Inventario".

Trae las columnas con los nombres que lee `pipeline/excel.py`, listas desplegables
para Estado y Moneda, y —si el loteo ya se construyó— una fila por parcela del KMZ
con lo que muestra hoy. Así nadie tiene que adivinar cómo se escribe "2-7", que es
lo que más confunde: un lote que no calza con el dibujo queda "no disponible".
"""
from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline.config import ESTADOS

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

COLUMNAS = ("Parcela", "Estado", "Precio", "Moneda", "Superficie m2", "Link de pago")
ANCHOS = (12, 16, 16, 10, 16, 44)
MONEDAS = ("CLP", "UF")
# Hasta dónde llegan las listas desplegables y el formato de texto de "Parcela".
FILAS_PREPARADAS = 2000

# Sin loteo construido no se sabe qué parcelas hay: tres filas de muestra.
EJEMPLO = (
    ("1", "Disponible", 9990000, "CLP", 5000, None),
    ("2", "Reservado", 10490000, "CLP", 5200, None),
    ("3", "Vendido", 395, "UF", 5000, None),
)

INSTRUCCIONES = (
    ("Cómo llenar el inventario", None),
    ("Una fila por parcela. Guarda el archivo y súbelo en \"Inventario\".", None),
    (None, None),
    ("Columna", "Qué va"),
    ("Parcela", ("El número del lote tal como está en el plano (KMZ). Con varias etapas, "
                 "etapa-número: 2-7 es la parcela 7 de la etapa 2. Es la única obligatoria.")),
    ("Estado", ("Disponible, Reservado, Vendido, No disponible o No en venta. "
                "Vacío cuenta como No disponible.")),
    ("Precio", "Solo el número, sin $ ni puntos. Vacío o 0: la ficha dice \"A consultar\"."),
    ("Moneda", "CLP o UF. Vacío es CLP."),
    ("Superficie m2", "En metros cuadrados, la de la escritura. Vacío: la ficha no la muestra."),
    ("Link de pago", "Opcional. Si lo pones, la ficha muestra el botón para pagar."),
    (None, None),
    ("Las parcelas que no estén en la planilla quedan como No disponible.", None),
)


def plantilla(parcelas: list[dict] | None = None) -> bytes:
    """El .xlsx listo para llenar; con `parcelas` (las de `parcelas.json`), prellenado."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Inventario"
    _encabezado(hoja)
    for fila in (_filas_de(parcelas) if parcelas else EJEMPLO):
        hoja.append(fila)
    _preparar_columnas(hoja)
    _instrucciones(libro.create_sheet("Cómo llenarla"))

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _filas_de(parcelas: list[dict]) -> list[tuple]:
    etiquetas = {clave: datos["etiqueta"] for clave, datos in ESTADOS.items()}
    return [(
        str(p["id"]),
        etiquetas.get(p.get("estado"), ""),
        p.get("precio"),
        p.get("moneda") or "CLP",
        # Sin planilla, la superficie que trae parcelas.json es la del dibujo: al
        # prellenarla pasaría por la oficial sin que nadie la haya escrito.
        p.get("superficie_m2") if p.get("en_planilla") else None,
        p.get("link_pago"),
    ) for p in parcelas]


def _encabezado(hoja) -> None:
    hoja.append(COLUMNAS)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="18181B")
        celda.alignment = Alignment(vertical="center")
    hoja.freeze_panes = "A2"


def _preparar_columnas(hoja) -> None:
    for posicion, ancho in enumerate(ANCHOS, start=1):
        hoja.column_dimensions[openpyxl.utils.get_column_letter(posicion)].width = ancho
    # "2-7" escrito en una celda General, Excel lo convierte en fecha (2 de julio):
    # la columna Parcela va como texto desde antes que alguien escriba en ella.
    for (celda,) in hoja.iter_rows(min_row=2, max_row=FILAS_PREPARADAS, max_col=1):
        celda.number_format = "@"

    estados = DataValidation(type="list", allow_blank=True, showErrorMessage=True,
                             formula1='"' + ",".join(e["etiqueta"] for e in ESTADOS.values()) + '"',
                             errorTitle="Estado", error="Elige un estado de la lista.")
    monedas = DataValidation(type="list", allow_blank=True, showErrorMessage=True,
                             formula1='"' + ",".join(MONEDAS) + '"',
                             errorTitle="Moneda", error="CLP o UF.")
    estados.add(f"B2:B{FILAS_PREPARADAS}")
    monedas.add(f"D2:D{FILAS_PREPARADAS}")
    hoja.add_data_validation(estados)
    hoja.add_data_validation(monedas)


def _instrucciones(hoja) -> None:
    for fila in INSTRUCCIONES:
        hoja.append(fila)
    hoja["A1"].font = Font(bold=True, size=14)
    for celda in hoja[4]:
        celda.font = Font(bold=True)
    hoja.column_dimensions["A"].width = 16
    hoja.column_dimensions["B"].width = 90
    for (_, texto) in hoja.iter_rows(min_row=5, max_row=10):
        texto.alignment = Alignment(wrap_text=True, vertical="top")


__all__ = ["MIME_XLSX", "plantilla"]
